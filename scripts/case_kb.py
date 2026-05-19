#!/usr/bin/env python3
"""Build and query a local case knowledge base (paragraph-level retrieval).

Usage examples:
  python scripts/case_kb.py index --source "/T20050026/deer-flow/data/case-kb" --db "/T20050026/deer-flow/data/case-kb/index.sqlite"
  python scripts/case_kb.py search --db "/T20050026/deer-flow/data/case-kb/index.sqlite" --query "正泰集团 数字化转型 组织变革" --top-k 8
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import re
import shutil
import sqlite3
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from textwrap import wrap


# Let scripts import harness utilities without installation.
REPO_ROOT = Path(__file__).resolve().parent.parent
HARNESS_ROOT = REPO_ROOT / "backend" / "packages" / "harness"
if str(HARNESS_ROOT) not in sys.path:
    sys.path.insert(0, str(HARNESS_ROOT))


@dataclass
class Chunk:
    source_path: str
    source_name: str
    chunk_index: int
    text: str


def _iter_source_files(source_dir: Path) -> list[Path]:
    files: list[Path] = []
    for ext in ("*.pdf", "*.md", "*.txt", "*.docx", "*.doc", "*.pptx", "*.ppt", "*.xlsx", "*.xls"):
        files.extend(source_dir.glob(ext))
    return sorted(set(files))


def _is_supported_file(path: Path) -> bool:
    return path.suffix.lower() in {".pdf", ".md", ".txt", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls"}


def _to_markdown_text(path: Path) -> str | None:
    """Convert source file into markdown/plain text.

    Strategy:
    1) .md/.txt read directly
    2) if sibling .md exists, use it
    3) use deerflow converter (preferred)
    4) fallback to MarkItDown
    """
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        try:
            return path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return None

    sibling_md = path.with_suffix(".md")
    if sibling_md.exists():
        try:
            return sibling_md.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            pass

    # Preferred: existing deerflow conversion utility.
    try:
        from deerflow.utils.file_conversion import convert_file_to_markdown

        md_path = asyncio.run(convert_file_to_markdown(path))
        if md_path and Path(md_path).exists():
            return Path(md_path).read_text(encoding="utf-8", errors="ignore")
    except Exception:
        pass

    # Fallback: MarkItDown direct conversion.
    try:
        from markitdown import MarkItDown

        converter = MarkItDown()
        result = converter.convert(str(path))
        return result.text_content or ""
    except Exception:
        return None


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_paragraphs(text: str) -> list[str]:
    text = _normalize_text(text)
    if not text:
        return []
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    # Merge too-short paragraphs to reduce fragment noise.
    merged: list[str] = []
    buf = ""
    for p in paras:
        if len(p) < 80:
            buf = f"{buf} {p}".strip()
            continue
        if buf:
            merged.append(f"{buf} {p}".strip())
            buf = ""
        else:
            merged.append(p)
    if buf:
        merged.append(buf)
    return merged


def _build_chunks(source_path: Path, text: str, max_chars: int = 1200, overlap_chars: int = 180) -> list[Chunk]:
    paras = _split_paragraphs(text)
    chunks: list[Chunk] = []
    if not paras:
        return chunks

    cur = ""
    idx = 0
    for p in paras:
        candidate = f"{cur}\n\n{p}".strip() if cur else p
        if len(candidate) <= max_chars:
            cur = candidate
            continue
        if cur:
            chunks.append(
                Chunk(
                    source_path=str(source_path),
                    source_name=source_path.name,
                    chunk_index=idx,
                    text=cur,
                )
            )
            idx += 1
            tail = cur[-overlap_chars:] if overlap_chars > 0 else ""
            cur = f"{tail}\n\n{p}".strip() if tail else p
        else:
            chunks.append(
                Chunk(
                    source_path=str(source_path),
                    source_name=source_path.name,
                    chunk_index=idx,
                    text=p[:max_chars],
                )
            )
            idx += 1
            cur = p[max_chars - overlap_chars :] if len(p) > max_chars else ""

    if cur.strip():
        chunks.append(
            Chunk(
                source_path=str(source_path),
                source_name=source_path.name,
                chunk_index=idx,
                text=cur.strip(),
            )
        )

    return chunks


def _connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    return conn


def _init_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA journal_mode=WAL;
        PRAGMA synchronous=NORMAL;

        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_path TEXT UNIQUE NOT NULL,
            source_name TEXT NOT NULL,
            source_hash TEXT NOT NULL,
            chunk_count INTEGER NOT NULL DEFAULT 0,
            indexed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER NOT NULL,
            chunk_index INTEGER NOT NULL,
            text TEXT NOT NULL,
            FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
        );

        CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(
            text,
            source_name,
            source_path,
            content='',
            tokenize='unicode61 remove_diacritics 2'
        );
        """
    )
    conn.commit()


def _reindex_document(conn: sqlite3.Connection, path: Path, text: str) -> tuple[int, int]:
    source_path = str(path)
    source_name = path.name
    source_hash = hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()

    row = conn.execute("SELECT id, source_hash FROM documents WHERE source_path = ?", (source_path,)).fetchone()
    if row and row["source_hash"] == source_hash:
        return (0, 1)  # unchanged, skipped

    if row:
        doc_id = row["id"]
        conn.execute("DELETE FROM chunks WHERE document_id = ?", (doc_id,))
        conn.execute("DELETE FROM chunks_fts WHERE source_path = ?", (source_path,))
    else:
        conn.execute(
            "INSERT INTO documents(source_path, source_name, source_hash, chunk_count) VALUES (?, ?, ?, 0)",
            (source_path, source_name, source_hash),
        )
        doc_id = conn.execute("SELECT id FROM documents WHERE source_path = ?", (source_path,)).fetchone()["id"]

    chunks = _build_chunks(path, text)
    for chunk in chunks:
        conn.execute(
            "INSERT INTO chunks(document_id, chunk_index, text) VALUES (?, ?, ?)",
            (doc_id, chunk.chunk_index, chunk.text),
        )
        conn.execute(
            "INSERT INTO chunks_fts(text, source_name, source_path) VALUES (?, ?, ?)",
            (chunk.text, source_name, source_path),
        )

    conn.execute(
        """
        UPDATE documents
        SET source_hash = ?, chunk_count = ?, indexed_at = CURRENT_TIMESTAMP
        WHERE id = ?
        """,
        (source_hash, len(chunks), doc_id),
    )
    conn.commit()
    return (len(chunks), 0)


def cmd_index(args: argparse.Namespace) -> int:
    source_dir = Path(args.source).resolve()
    db_path = Path(args.db).resolve()
    if not source_dir.exists() or not source_dir.is_dir():
        print(f"[ERROR] source directory not found: {source_dir}")
        return 1

    conn = _connect(db_path)
    _init_schema(conn)

    files = _iter_source_files(source_dir)
    if not files:
        print(f"[WARN] no supported files found under: {source_dir}")
        return 0

    indexed_chunks = 0
    skipped_docs = 0
    failed_docs = 0

    for file_path in files:
        text = _to_markdown_text(file_path)
        if not text or not text.strip():
            failed_docs += 1
            print(f"[WARN] failed to parse: {file_path.name}")
            continue
        chunks_added, skipped = _reindex_document(conn, file_path, text)
        indexed_chunks += chunks_added
        skipped_docs += skipped
        status = "skipped" if skipped else f"indexed {chunks_added} chunks"
        print(f"[OK] {file_path.name}: {status}")

    doc_count = conn.execute("SELECT COUNT(*) AS c FROM documents").fetchone()["c"]
    chunk_count = conn.execute("SELECT COUNT(*) AS c FROM chunks").fetchone()["c"]
    print(
        f"\n[DONE] docs={doc_count}, chunks={chunk_count}, newly_indexed_chunks={indexed_chunks}, "
        f"skipped_docs={skipped_docs}, failed_docs={failed_docs}, db={db_path}"
    )
    return 0


def cmd_search(args: argparse.Namespace) -> int:
    db_path = Path(args.db).resolve()
    if not db_path.exists():
        print(f"[ERROR] db not found: {db_path}")
        return 1

    query = args.query.strip()
    if not query:
        print("[ERROR] query cannot be empty")
        return 1

    conn = _connect(db_path)
    sql = """
        SELECT
            c.id,
            c.text,
            d.source_name,
            d.source_path,
            c.chunk_index,
            bm25(chunks_fts) AS score
        FROM chunks_fts
        JOIN chunks c ON c.rowid = chunks_fts.rowid
        JOIN documents d ON d.id = c.document_id
        WHERE chunks_fts MATCH ?
        ORDER BY score
        LIMIT ?
    """
    rows = conn.execute(sql, (query, args.top_k)).fetchall()
    if not rows:
        print("[INFO] no hits")
        return 0

    for i, row in enumerate(rows, 1):
        snippet = row["text"][: args.snippet_chars].replace("\n", " ")
        print(f"[{i}] score={row['score']:.4f} file={row['source_name']} chunk={row['chunk_index']}")
        print(f"    path: {row['source_path']}")
        print(f"    text: {snippet}")
        print()
    return 0


def cmd_ingest(args: argparse.Namespace) -> int:
    """Ingest uploaded files to KB dir, then build/update index.

    Default behavior copies files from uploads to kb directory and keeps uploads intact.
    Use --move to move files instead of copy.
    """
    uploads_dir = Path(args.uploads).resolve()
    kb_dir = Path(args.kb_dir).resolve()
    db_path = Path(args.db).resolve()

    if not uploads_dir.exists() or not uploads_dir.is_dir():
        print(f"[ERROR] uploads directory not found: {uploads_dir}")
        return 1

    kb_dir.mkdir(parents=True, exist_ok=True)
    upload_files = sorted(p for p in uploads_dir.iterdir() if p.is_file() and _is_supported_file(p))
    if not upload_files:
        print(f"[WARN] no supported files found in uploads: {uploads_dir}")
        return 0

    copied = 0
    skipped = 0
    failed = 0
    action = "move" if args.move else "copy"

    for src in upload_files:
        dst = kb_dir / src.name
        try:
            # Skip identical file quickly by size + mtime heuristic if destination exists.
            if dst.exists() and dst.stat().st_size == src.stat().st_size and int(dst.stat().st_mtime) == int(src.stat().st_mtime):
                skipped += 1
                print(f"[SKIP] {src.name}: unchanged in kb dir")
                continue
            if args.move:
                shutil.move(str(src), str(dst))
            else:
                shutil.copy2(str(src), str(dst))
            copied += 1
            print(f"[OK] {action} {src.name} -> {dst}")
        except Exception as e:
            failed += 1
            print(f"[WARN] failed to {action} {src.name}: {e}")

    print(f"\n[INGEST] copied_or_moved={copied}, skipped={skipped}, failed={failed}")

    # Reuse index command directly
    index_args = argparse.Namespace(source=str(kb_dir), db=str(db_path))
    return cmd_index(index_args)


def _build_default_export_path(report_dir: Path, export_format: str, ts: str) -> Path:
    suffix_map = {
        "word": ".docx",
        "pdf": ".pdf",
        "md": ".md",
    }
    suffix = suffix_map.get(export_format, ".txt")
    return report_dir / f"case-analysis-export-{ts}{suffix}"


def _export_word(content: str, title: str, export_path: Path) -> str:
    from docx import Document  # type: ignore

    doc = Document()
    doc.add_heading(title, level=1)
    for line in content.splitlines():
        doc.add_paragraph(line)
    doc.save(export_path)
    return str(export_path)


def _export_pdf(content: str, export_path: Path) -> str:
    lines = content.splitlines()
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfgen import canvas

        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        c = canvas.Canvas(str(export_path), pagesize=A4)
        page_width, page_height = A4
        left = 36
        top = page_height - 40
        bottom = 40
        line_height = 15
        c.setFont("STSong-Light", 11)

        y = top
        for line in lines:
            wrapped_lines = wrap(line, width=54) or [""]
            for wrapped in wrapped_lines:
                if y < bottom:
                    c.showPage()
                    c.setFont("STSong-Light", 11)
                    y = top
                c.drawString(left, y, wrapped)
                y -= line_height
        c.save()
        return str(export_path)
    except Exception:
        from matplotlib import pyplot as plt
        from matplotlib.backends.backend_pdf import PdfPages

        page_size = 38
        with PdfPages(export_path) as pdf:
            for i in range(0, len(lines), page_size):
                fig = plt.figure(figsize=(8.27, 11.69))
                ax = fig.add_subplot(111)
                ax.axis("off")
                chunk = lines[i : i + page_size]
                wrapped_lines: list[str] = []
                for line in chunk:
                    wrapped_lines.extend(wrap(line, width=68) or [""])
                text_block = "\n".join(wrapped_lines)
                ax.text(0.02, 0.98, text_block, ha="left", va="top", fontsize=10, family="sans-serif")
                pdf.savefig(fig, bbox_inches="tight")
                plt.close(fig)
        return str(export_path)


def _export_md(content: str, export_path: Path) -> str:
    export_path.write_text(content if content.endswith("\n") else content + "\n", encoding="utf-8")
    return str(export_path)


def export_case_content(content: str, title: str, export_format: str, export_path: str) -> str:
    report_dir = REPO_ROOT / "skills" / "public" / "graduate-case-analysis" / "exports"
    report_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    fmt = export_format.lower()
    if fmt in {"doc", "docx", "auto"}:
        fmt = "word"

    if fmt in {"word", "pdf", "md"}:
        file_path = Path(export_path).resolve() if export_path else _build_default_export_path(report_dir, fmt, ts)
        if fmt == "word":
            return _export_word(content, title, file_path)
        if fmt == "pdf":
            return _export_pdf(content, file_path)
        return _export_md(content, file_path)
    raise ValueError(f"unsupported export format: {export_format}")


def cmd_export(args: argparse.Namespace) -> int:
    if args.export_format == "none":
        print("[INFO] export skipped: default behavior is web output only (no file exported)")
        return 0

    content_path = Path(args.content_file).resolve()
    if not content_path.exists() or not content_path.is_file():
        print(f"[ERROR] content file not found: {content_path}")
        return 1
    content = content_path.read_text(encoding="utf-8", errors="ignore").strip()
    if not content:
        print(f"[ERROR] content file is empty: {content_path}")
        return 1

    try:
        file_path = export_case_content(
            content=content,
            title=args.title,
            export_format=args.export_format,
            export_path=args.export_path,
        )
    except Exception as exc:
        print(f"[ERROR] export failed: {exc}")
        return 1

    print("[OK] export done")
    print(f"  source(latest_response): {content_path}")
    print(f"  requested: {args.export_format}")
    print(f"  file: {file_path}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Local case knowledge-base index/search.")
    sub = parser.add_subparsers(dest="command", required=True)

    p_index = sub.add_parser("index", help="Build or update KB index")
    p_index.add_argument("--source", required=True, help="Source directory containing PDFs/docs")
    p_index.add_argument("--db", required=True, help="SQLite database path")
    p_index.set_defaults(func=cmd_index)

    p_search = sub.add_parser("search", help="Search indexed paragraphs")
    p_search.add_argument("--db", required=True, help="SQLite database path")
    p_search.add_argument("--query", required=True, help="Search query")
    p_search.add_argument("--top-k", type=int, default=8, help="Number of hits to return")
    p_search.add_argument("--snippet-chars", type=int, default=220, help="Snippet length for printing")
    p_search.set_defaults(func=cmd_search)

    p_ingest = sub.add_parser("ingest", help="Ingest uploads into KB dir, then build/update index")
    p_ingest.add_argument("--uploads", default="/mnt/user-data/uploads", help="Uploads directory")
    p_ingest.add_argument("--kb-dir", default=str(REPO_ROOT / "data" / "case-kb"), help="KB storage directory")
    p_ingest.add_argument("--db", default=str(REPO_ROOT / "data" / "case-kb" / "index.sqlite"), help="SQLite index path")
    p_ingest.add_argument("--move", action="store_true", help="Move files from uploads instead of copying")
    p_ingest.set_defaults(func=cmd_ingest)

    p_export = sub.add_parser("export", help="Export latest case-analysis response to document")
    p_export.add_argument("--content-file", required=True, help="Path to latest response text/markdown content")
    p_export.add_argument(
        "--export-format",
        choices=("none", "auto", "word", "doc", "docx", "pdf", "md"),
        default="none",
        help="Document format. Default is none (web output only); auto/doc/docx map to word.",
    )
    p_export.add_argument("--export-path", default="", help="Optional output document path")
    p_export.add_argument("--title", default="案例分析生成结果", help="Document title for Word export")
    p_export.set_defaults(func=cmd_export)

    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())

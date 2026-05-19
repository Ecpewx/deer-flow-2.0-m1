"""
试卷 / 论文选题导出工具
将 Markdown 文件转换为 .docx 或 .pdf 格式，支持 LaTeX 公式渲染（图片方式）。

用法：
  python export.py --input exam.md --format docx --output exam.docx
  python export.py --input topics.md --format pdf --output topics.pdf
  python export.py --input exam.md --format both --output-dir /mnt/user-data/workspace/

依赖：python-docx, markdown, weasyprint (PDF), matplotlib (公式渲染)
"""

import argparse
import base64
import io
import logging
import os
import re
import subprocess
import sys

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# 自动安装缺失依赖
# ---------------------------------------------------------------------------
REQUIRED_PACKAGES = {
    "docx": "python-docx",
    "markdown": "markdown",
    "matplotlib": "matplotlib",
}

PDF_PACKAGES = {
    "weasyprint": "weasyprint",
}


def _ensure_packages(packages: dict) -> None:
    for module, pip_name in packages.items():
        try:
            __import__(module)
        except ImportError:
            logger.info("正在安装 %s ...", pip_name)
            # 先尝试 uv pip（更可靠），失败再退回普通 pip
            installed = False
            try:
                subprocess.check_call(
                    ["uv", "pip", "install", pip_name, "-q"],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
                installed = True
            except (subprocess.CalledProcessError, FileNotFoundError):
                pass
            if not installed:
                subprocess.check_call(
                    [sys.executable, "-m", "pip", "install", pip_name, "-q"],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )


# ---------------------------------------------------------------------------
# LaTeX 公式 → PNG（matplotlib 渲染）
# ---------------------------------------------------------------------------
def _render_formula_png(latex: str, display: bool = False, dpi: int = 150):
    """将 LaTeX 公式渲染为 PNG 字节。失败时返回 None。"""
    _ensure_packages({"matplotlib": "matplotlib"})
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        fontsize = 16 if display else 13
        formula_str = "$" + latex + "$"

        fig = plt.figure(figsize=(0.1, 0.1))
        fig.patch.set_alpha(0)
        t = fig.text(0.5, 0.5, formula_str, fontsize=fontsize,
                     ha="center", va="center")
        fig.canvas.draw()
        renderer = fig.canvas.get_renderer()
        bbox = t.get_window_extent(renderer=renderer)

        pad = 12
        w = (bbox.width + pad * 2) / dpi
        h = (bbox.height + pad) / dpi
        fig.set_size_inches(max(w, 0.3), max(h, 0.2))

        buf = io.BytesIO()
        fig.savefig(buf, format="png", bbox_inches="tight",
                    pad_inches=0.04, transparent=True, dpi=dpi)
        plt.close(fig)
        buf.seek(0)
        return buf.read()
    except Exception as exc:
        logger.warning("公式渲染失败（%s）: %s", latex, exc)
        return None


# ---------------------------------------------------------------------------
# 行内公式处理工具
# ---------------------------------------------------------------------------
_DISPLAY_RE = re.compile(r"\$\$(.*?)\$\$", re.DOTALL)
_INLINE_RE = re.compile(r"(?<!\$)\$(?!\$)(.*?)(?<!\$)\$(?!\$)")


def _split_inline_formulas(text: str):
    """将含有 $...$ 的文本拆成片段，返回 [(content, is_formula), ...]"""
    parts = []
    last = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > last:
            parts.append((text[last:m.start()], False))
        parts.append((m.group(1), True))
        last = m.end()
    if last < len(text):
        parts.append((text[last:], False))
    return parts


# ---------------------------------------------------------------------------
# Markdown → HTML（公式转为 base64 PNG img 标签）
# ---------------------------------------------------------------------------
def _formulas_to_img_tags(md_text: str) -> str:
    """将 $$...$$ 和 $...$ 替换为 <img> 标签（base64 PNG）。"""

    def _replace_display(m):
        png = _render_formula_png(m.group(1).strip(), display=True)
        if png is None:
            return m.group(0)
        b64 = base64.b64encode(png).decode()
        return (
            '<div style="text-align:center;margin:0.6em 0;">'
            '<img src="data:image/png;base64,' + b64 + '" '
            'style="vertical-align:middle;max-height:2.5em;"/></div>'
        )

    def _replace_inline(m):
        png = _render_formula_png(m.group(1).strip(), display=False)
        if png is None:
            return m.group(0)
        b64 = base64.b64encode(png).decode()
        return (
            '<img src="data:image/png;base64,' + b64 + '" '
            'style="vertical-align:middle;max-height:1.6em;"/>'
        )

    md_text = _DISPLAY_RE.sub(_replace_display, md_text)
    md_text = _INLINE_RE.sub(_replace_inline, md_text)
    return md_text


def _md_to_html(md_text: str) -> str:
    """Markdown -> HTML，公式渲染为嵌入图片。"""
    _ensure_packages({"markdown": "markdown"})
    import markdown

    processed = _formulas_to_img_tags(md_text)
    html_body = markdown.markdown(
        processed,
        extensions=["tables", "fenced_code", "toc", "nl2br"],
    )

    css = """
body { font-family: "SimSun", serif; max-width: 800px; margin: 40px auto;
       padding: 0 20px; line-height: 1.8; color: #333; }
h1 { font-size: 22pt; text-align: center; margin-bottom: 0.3em; }
h2 { font-size: 16pt; border-bottom: 1px solid #ccc; padding-bottom: 4px; margin-top: 1.5em; }
h3 { font-size: 13pt; }
table { border-collapse: collapse; width: 100%; margin: 1em 0; }
th, td { border: 1px solid #999; padding: 6px 10px; text-align: left; }
th { background: #f0f0f0; }
blockquote { border-left: 3px solid #aaa; padding-left: 12px; color: #666; }
code { background: #f5f5f5; padding: 2px 4px; border-radius: 3px; }
"""
    return (
        '<!DOCTYPE html>\n<html lang="zh-CN">\n<head>\n<meta charset="utf-8">\n'
        "<style>" + css + "</style>\n</head>\n<body>\n" +
        html_body + "\n</body>\n</html>"
    )


# ---------------------------------------------------------------------------
# 导出为 DOCX
# ---------------------------------------------------------------------------
def _add_paragraph_with_formulas(doc, text: str):
    """向 doc 添加段落，$...$ 公式渲染为内联图片。"""
    from docx.shared import Pt

    p = doc.add_paragraph()
    for chunk, is_formula in _split_inline_formulas(text):
        if not chunk:
            continue
        if is_formula:
            png = _render_formula_png(chunk.strip(), display=False)
            if png:
                run = p.add_run()
                run.add_picture(io.BytesIO(png), height=Pt(14))
            else:
                p.add_run("$" + chunk + "$")
        else:
            bold_parts = re.split(r"\*\*(.*?)\*\*", chunk)
            for idx, bp in enumerate(bold_parts):
                if not bp:
                    continue
                run = p.add_run(bp)
                run.bold = (idx % 2 == 1)
    return p


def _find_template(output_path: str) -> str | None:
    """根据输出文件名自动匹配格式模板路径，找不到时返回 None。"""
    templates_dir = os.path.join(os.path.dirname(__file__), "..", "templates")
    templates_dir = os.path.abspath(templates_dir)
    basename = os.path.basename(output_path)
    if "答案" in basename:
        candidate = os.path.join(templates_dir, "试卷答案及评分标准模板.docx")
    else:
        candidate = os.path.join(templates_dir, "浙江财经大学出卷模板.docx")
    return candidate if os.path.exists(candidate) else None


def export_docx(md_text: str, output_path: str) -> str:
    """将 Markdown 文本导出为 .docx 文件，公式渲染为图片。"""
    _ensure_packages(REQUIRED_PACKAGES)

    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    template_path = _find_template(output_path)
    if template_path:
        logger.info("使用模板: %s", template_path)
        try:
            import io as _io
            with open(template_path, "rb") as _f:
                template_bytes = _io.BytesIO(_f.read())
            doc = Document(template_bytes)
            # 清除模板中的所有现有正文段落和表格，保留样式
            body = doc.element.body
            for child in list(body):
                tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                if tag in ("p", "tbl", "sdt"):
                    body.remove(child)
        except Exception as e:
            logger.warning(
                "模板加载失败（%s），可能是旧版 .doc 格式，请用 Word/WPS 另存为 .docx 格式后替换。"
                "当前使用空白文档继续。错误: %s", template_path, e
            )
            template_path = None
    if not template_path:
        logger.warning("未找到匹配的模板，使用空白文档")
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "宋体"
        style.font.size = Pt(12)
        for section in doc.sections:
            section.top_margin = Cm(2.54)
            section.bottom_margin = Cm(2.54)
            section.left_margin = Cm(3.17)
            section.right_margin = Cm(3.17)

    lines = md_text.split("\n")
    in_table = False
    table_rows = []
    in_code_block = False

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("```"):
            in_code_block = not in_code_block
            continue
        if in_code_block:
            p = doc.add_paragraph(line)
            run = p.runs[0] if p.runs else p.add_run("")
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            continue

        if "|" in stripped and stripped.startswith("|"):
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            if all(re.match(r"^[-:]+$", c) for c in cells):
                continue
            table_rows.append(cells)
            in_table = True
            continue
        elif in_table:
            if table_rows:
                cols = len(table_rows[0])
                # "Table Grid" 在中文模板中可能不存在，优雅降级
                try:
                    tbl = doc.add_table(rows=len(table_rows), cols=cols, style="Table Grid")
                except KeyError:
                    tbl = doc.add_table(rows=len(table_rows), cols=cols)
                for i, row_data in enumerate(table_rows):
                    for j, cell_text in enumerate(row_data):
                        if j < cols:
                            tbl.rows[i].cells[j].text = cell_text
                            if i == 0:
                                for run in tbl.rows[i].cells[j].paragraphs[0].runs:
                                    run.bold = True
            table_rows = []
            in_table = False

        if not stripped:
            continue

        # Display math $$...$$
        dm = _DISPLAY_RE.fullmatch(stripped)
        if dm:
            png = _render_formula_png(dm.group(1).strip(), display=True)
            if png:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(io.BytesIO(png), height=Pt(20))
            else:
                p = doc.add_paragraph(stripped)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("# ").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.bold = True
            if level == 1:
                run.font.size = Pt(16)
            elif level == 2:
                run.font.size = Pt(14)
            else:
                run.font.size = Pt(12)
            continue

        if stripped.startswith(">"):
            text = stripped.lstrip("> ").strip()
            p = _add_paragraph_with_formulas(doc, text)
            p.paragraph_format.left_indent = Cm(1)
            for run in p.runs:
                run.font.size = Pt(10)
            continue

        if stripped in ("---", "***", "___"):
            doc.add_paragraph("─" * 50)
            continue

        text = re.sub(r"^\s*[-*]\s+", "• ", stripped)
        text = re.sub(r"`(.*?)`", r"\1", text)
        _add_paragraph_with_formulas(doc, text)

    if in_table and table_rows:
        cols = len(table_rows[0])
        tbl = doc.add_table(rows=len(table_rows), cols=cols, style="Table Grid")
        for i, row_data in enumerate(table_rows):
            for j, cell_text in enumerate(row_data):
                if j < cols:
                    tbl.rows[i].cells[j].text = cell_text

    doc.save(output_path)
    return "✅ DOCX 导出成功：" + output_path


# ---------------------------------------------------------------------------
# 导出为 PDF
# ---------------------------------------------------------------------------
def export_pdf(md_text: str, output_path: str) -> str:
    """将 Markdown 文本导出为 .pdf 文件，公式渲染为嵌入图片。"""
    _ensure_packages(REQUIRED_PACKAGES)
    _ensure_packages(PDF_PACKAGES)

    from weasyprint import HTML

    html_content = _md_to_html(md_text)
    HTML(string=html_content).write_pdf(output_path)
    return "✅ PDF 导出成功：" + output_path


# ---------------------------------------------------------------------------
# 主入口
# ---------------------------------------------------------------------------
def main() -> None:
    parser = argparse.ArgumentParser(description="试卷/论文选题 Markdown 导出工具")
    parser.add_argument("--input", required=True, help="输入 Markdown 文件路径")
    parser.add_argument(
        "--format",
        required=True,
        choices=["docx", "pdf", "both"],
        help="导出格式：docx / pdf / both",
    )
    parser.add_argument("--output", help="输出文件路径（单格式时使用）")
    parser.add_argument("--output-dir", help="输出目录（both 模式时使用）")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print("❌ 输入文件不存在：" + args.input)
        sys.exit(1)

    with open(args.input, "r", encoding="utf-8") as f:
        md_text = f.read()

    basename = os.path.splitext(os.path.basename(args.input))[0]

    if args.format == "docx":
        output = args.output or basename + ".docx"
        print(export_docx(md_text, output))
    elif args.format == "pdf":
        output = args.output or basename + ".pdf"
        print(export_pdf(md_text, output))
    elif args.format == "both":
        out_dir = args.output_dir or os.path.dirname(args.input) or "."
        os.makedirs(out_dir, exist_ok=True)
        print(export_docx(md_text, os.path.join(out_dir, basename + ".docx")))
        print(export_pdf(md_text, os.path.join(out_dir, basename + ".pdf")))


if __name__ == "__main__":
    main()
